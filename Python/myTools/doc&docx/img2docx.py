import sys
import os
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION


# 支持的图片格式
IMAGE_EXTS = {
    ".jpg", ".jpeg", ".png", ".bmp", ".gif",
    ".tif", ".tiff", ".webp"
}


def is_image_file(path: Path):
    return path.suffix.lower() in IMAGE_EXTS


def fit_size(img_width, img_height, max_width, max_height):
    """
    根据图片比例，将图片缩放到页面可用区域内
    """
    img_ratio = img_width / img_height
    page_ratio = max_width / max_height

    if img_ratio > page_ratio:
        width = max_width
        height = max_width / img_ratio
    else:
        height = max_height
        width = max_height * img_ratio

    return width, height


def convert_image_if_needed(image_path: Path, temp_dir: Path):
    """
    将部分 Word 不一定支持的图片格式转换为 PNG
    例如 webp
    """
    suffix = image_path.suffix.lower()

    if suffix in [".jpg", ".jpeg", ".png", ".bmp", ".gif", ".tif", ".tiff"]:
        return image_path

    try:
        img = Image.open(image_path)
        img = img.convert("RGB")

        output_path = temp_dir / f"{image_path.stem}.png"
        img.save(output_path, "PNG")

        return output_path
    except Exception as e:
        print(f"转换失败：{image_path}，原因：{e}")
        return None


def images_to_docx(folder_path):
    folder = Path(folder_path)

    if not folder.exists() or not folder.is_dir():
        print("请拖入一个有效的文件夹。")
        return

    images = [
        p for p in folder.iterdir()
        if p.is_file() and is_image_file(p)
    ]

    images.sort(key=lambda x: x.name.lower())

    if not images:
        print("文件夹内没有找到支持的图片。")
        return

    output_docx = folder / f"{folder.name}.docx"

    doc = Document()

    # 设置 A4 页面
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

    # 设置页边距
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    max_width = section.page_width - section.left_margin - section.right_margin
    max_height = section.page_height - section.top_margin - section.bottom_margin

    with tempfile.TemporaryDirectory() as tmp:
        temp_dir = Path(tmp)

        for index, image_path in enumerate(images):
            print(f"正在处理：{image_path.name}")

            usable_image_path = convert_image_if_needed(image_path, temp_dir)

            if usable_image_path is None:
                continue

            try:
                with Image.open(usable_image_path) as img:
                    img_width, img_height = img.size

                pic_width, pic_height = fit_size(
                    img_width,
                    img_height,
                    max_width,
                    max_height
                )

                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                run = paragraph.add_run()
                run.add_picture(
                    str(usable_image_path),
                    width=pic_width,
                    height=pic_height
                )

                # 每张图一页，最后一张不加分页
                if index != len(images) - 1:
                    doc.add_page_break()

            except Exception as e:
                print(f"添加失败：{image_path}，原因：{e}")

        doc.save(output_docx)

    print()
    print("转换完成！")
    print(f"输出文件：{output_docx}")


if __name__ == "__main__":
    if len(sys.argv) >= 2:
        folder_path = sys.argv[1]
    else:
        folder_path = input("请输入或拖入图片文件夹路径：").strip('"')

    images_to_docx(folder_path)

    input("按回车键退出...")
